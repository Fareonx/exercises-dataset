# -*- coding: utf-8 -*-
"""Builds plan/index.html, the artifact page and plan/qidalanma-plani.xlsx from data.py."""
import json, os, sys, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from data import *

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
OUT_PLAN = os.path.join(REPO, "plan")
# skeleton-less page for claude.ai Artifact publishing (optional)
OUT_ART = os.environ.get("ART_OUT", os.path.join(HERE, "..", "..", ".artifact-build"))
os.makedirs(OUT_PLAN, exist_ok=True)
os.makedirs(OUT_ART, exist_ok=True)

ex_db = {e["id"]: e for e in json.load(open(os.path.join(REPO, "data/exercises.json"), encoding="utf-8"))}

def lab(f):
    return f"{f[1]} ({f[2]})"

def fmt_qty(fid, g):
    if fid == "yumurta":
        n = round(g / 55)
        return f"≈ {n} ədəd"
    if fid in ("sud15", "sud35", "kefir"):
        return f"{g/1000:.2f}".rstrip("0").rstrip(".").replace(".", ",") + " l"
    if g >= 1000:
        return f"{g/1000:.1f}".replace(".", ",") + " kq"
    return f"{int(-(-g // 10) * 10)} q"

def shopping():
    tot = {}
    for _, _, meals in MENU:
        for _, _, items in meals:
            for fid, g in items:
                if fid in OUTSIDE: continue
                tot[fid] = tot.get(fid, 0) + g
    order = {f[0]: i for i, f in enumerate(FOODS)}
    rows = []
    for fid in sorted(tot, key=lambda k: order[k]):
        f = FOOD[fid]; g = tot[fid]
        rows.append(dict(id=fid, name=f[1], it=f[2], cat=f[3], grams=g, qty=fmt_qty(fid, g),
                         eur=round(g / 1000 * PRICE_EUR_KG.get(fid, 0), 2), eur_kg=PRICE_EUR_KG.get(fid, 0)))
    return rows

SHOP = shopping()

def plan_json():
    return dict(
        profile=PROFILE,
        meals=MEALS,
        foods=[dict(id=f[0], name=f[1], it=f[2], cat=f[3], k=f[4], p=f[5], c=f[6], f=f[7], hint=f[8]) for f in FOODS],
        menu=[dict(day=d, type=t, meals=[dict(meal=m, title=ti, items=[[fid, g] for fid, g in it]) for m, ti, it in ms]) for d, t, ms in MENU],
        timing={k: [v[0], [list(r) for r in v[1]]] for k, v in TIMING.items()},
        tennis=TENNIS_TIPS, rules=GENERAL_RULES, bologna=[list(x) for x in BOLOGNA_TIPS],
        week=[list(x) for x in WEEK_SCHEDULE], trainRules=TRAINING_RULES,
        workout={k: [dict(id=i, name=n, sets=s, rest=r, cue=c, gif=ex_db[i]["gif_url"], en=ex_db[i]["name"]) for i, n, s, r, c in v] for k, v in WORKOUT.items()},
        shopping=[dict(name=s["name"], it=s["it"], qty=s["qty"], eur=s["eur"]) for s in SHOP],
    )

def build_html():
    tpl = open(os.path.join(HERE, "template.html"), encoding="utf-8").read()
    data = json.dumps(plan_json(), ensure_ascii=False).replace("</", "<\\/")
    page = tpl.replace("/*__PLAN__*/null", data)
    open(os.path.join(OUT_ART, "index.html"), "w", encoding="utf-8").write(page)
    full = ('<!doctype html>\n<html lang="az">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">\n'
            '<style>[hidden]{display:none!important}img{max-width:100%}</style>\n</head>\n<body>\n'
            + page + '\n</body>\n</html>\n')
    open(os.path.join(OUT_PLAN, "index.html"), "w", encoding="utf-8").write(full)

# ------------------------------------------------------------------ Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

F = "Arial"
H1 = Font(name=F, size=14, bold=True)
HB = Font(name=F, size=10, bold=True, color="FFFFFF")
B = Font(name=F, size=10, bold=True)
N = Font(name=F, size=10)
INP = Font(name=F, size=10, color="0000FF")
LINK = Font(name=F, size=10, color="008000")
MUT = Font(name=F, size=9, italic=True, color="5B6760")
HFILL = PatternFill("solid", fgColor="1D6B4C")
YFILL = PatternFill("solid", fgColor="FFF59D")
SFILL = PatternFill("solid", fgColor="E8ECE6")
thin = Side(style="thin", color="D3D9D1")
BOX = Border(bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, row, labels, widths=None):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=l); c.font = HB; c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w

def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=c.font.size if c.font and c.font.size else 10, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False, color=c.font.color if c.font else None)

START = datetime.date(2026, 10, 9)   # tomorrow: first gym day
NDAYS = 120
LOGROWS = 2000

def build_xlsx():
    wb = Workbook()
    # ---------------- Təlimat
    ws = wb.active; ws.title = "Təlimat"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Kütlə planı · Bolonya — necə istifadə etmək", H1),
        ("", N),
        ("1. “Gündəlik” vərəqində hər yediyini bir sətirə yaz: Tarix, Yemək, Məhsul (siyahıdan seç), Qram. kkal və makrolar avtomatik hesablanır.", N),
        ("2. “Günlər” vərəqində həmin günün tipini seç (Zal / İstirahət) və tennis oynamısansa dəqiqəni yaz. Norma avtomatik dəyişir.", N),
        ("3. Səhər çəkisini “Günlər” vərəqinin “Çəki” sütununa yaz (həftədə ən azı 4 dəfə).", N),
        ("4. “Həftəlik” vərəqi orta kalorini, proteini və çəki dəyişməsini göstərir, normanı artırmaq/azaltmaq lazımdırsa deyir.", N),
        ("5. Çəki dəyişəndə “Profil” vərəqində çəkini yenilə — bütün normalar yenidən hesablanır.", N),
        ("", N),
        ("Rənglər: göy rəqəm = sənin dəyişə biləcəyin giriş xanası; sarı fon = doldurmalı olduğun xana; qara = formula (toxunma).", B),
        ("“Gündəlik” vərəqindəki ilk sətirlər nümunədir (9 oktyabr, Cümə menyusu). Yediyinə görə dəyiş və ya sil.", N),
        ("Öz məhsulunu “Qida bazası” vərəqinin sonundakı boş sarı sətirlərə əlavə et (qablaşdırmadakı “per 100 g” dəyərləri).", N),
        ("", N),
        ("Qida dəyərləri: CREA (İtaliya qida cədvəlləri), USDA FoodData Central və İtaliya supermarket etiketlərinə əsaslanan orta rəqəmlərdir; məhsuldan məhsula ±10% fərq ola bilər.", MUT),
        ("Qiymətlər Bolonya supermarketləri (Lidl, Conad, Coop, Esselunga) üçün təxminidir, 2026.", MUT),
        ("Bu plan tibbi məsləhət deyil. Sağlamlıq problemin varsa, həkimlə məsləhətləş.", MUT),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True)

    # ---------------- Profil
    ws = wb.create_sheet("Profil")
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 60
    ws["A1"] = "Profil və hesablama"; ws["A1"].font = H1
    inputs = [
        ("Çəki, kq", PROFILE["weight"], "Səhər, ac qarına"),
        ("Boy, sm", PROFILE["height"], ""),
        ("Yaş", PROFILE["age"], "Sənin dediyin: 18–24 yaş; 21 götürülüb — dəqiq yaşını yaz"),
        ("Aktivlik əmsalı (məşqsiz)", PROFILE["activity"], "1,3 çox az hərəkət · 1,4 tələbə, piyada gəzir · 1,5 ayaq üstə iş"),
        ("Kütlə üçün artıq kalori, kkal", PROFILE["surplus"], "Həftəlik çəki artımına görə ±150 dəyiş"),
        ("Zal məşqi (≈60 dəq), kkal", PROFILE["gym_kcal"], "Güc məşqi ≈ (MET 5 − 1) × çəki × saat"),
        ("Protein, q/kq", PROFILE["protein_gkg"], "Əzələ artımı üçün 1,6–2,2"),
        ("Yağ, kalorinin payı", PROFILE["fat_pct"], "20–30%"),
    ]
    ws["A3"] = "Giriş"; ws["A3"].font = B
    for i, (l, v, n) in enumerate(inputs):
        r = 4 + i
        ws.cell(row=r, column=1, value=l).font = N
        c = ws.cell(row=r, column=2, value=v); c.font = INP; c.fill = YFILL
        ws.cell(row=r, column=3, value=n).font = MUT
    ws["B11"].number_format = "0%"
    ws["A13"] = "Nəticə"; ws["A13"].font = B
    calc = [
        ("Bazal metabolizm (BMR), kkal", "=ROUND(10*B4+6.25*B5-5*B6+5,0)", "Mifflin–St Jeor, kişi"),
        ("Gündəlik xərc, məşqsiz, kkal", "=ROUND((10*B4+6.25*B5-5*B6+5)*B7,0)", "BMR × aktivlik"),
        ("İstirahət günü norması, kkal", "=ROUND(((10*B4+6.25*B5-5*B6+5)*B7+B8)/10,0)*10", "xərc + artıq"),
        ("Zal günü norması, kkal", "=ROUND(((10*B4+6.25*B5-5*B6+5)*B7+B8+B9)/10,0)*10", "xərc + artıq + məşq"),
        ("Tennis tək, 60 dəq, kkal", "=ROUND(7*B4,0)", "(MET 8 − 1) × çəki"),
        ("Tennis cüt, 60 dəq, kkal", "=ROUND(5*B4,0)", "(MET 6 − 1) × çəki"),
        ("Protein, minimum q/gün", "=ROUND(B10*B4,0)", ""),
        ("Yağ, istirahət / zal, q", '=ROUND(B16*B11/9,0)&" / "&ROUND(B17*B11/9,0)', ""),
        ("Karbohidrat, istirahət, q", "=ROUND((B16-B20*4-ROUND(B16*B11/9,0)*9)/4,0)", "qalan kalori"),
        ("Karbohidrat, zal, q", "=ROUND((B17-B20*4-ROUND(B17*B11/9,0)*9)/4,0)", ""),
        ("Su, l/gün", "2,5–3", "tennis/zal günü +0,5–1 l"),
        ("Hədəf çəki artımı, kq/həftə", "0,2–0,3", "təmiz kütlə"),
    ]
    for i, (l, f, n) in enumerate(calc):
        r = 14 + i
        ws.cell(row=r, column=1, value=l).font = B if "norması" in l else N
        c = ws.cell(row=r, column=2, value=f); c.font = B if "norması" in l else N
        if isinstance(f, str) and f.startswith("="): c.number_format = "#,##0"
        ws.cell(row=r, column=3, value=n).font = MUT

    # ---------------- Qida bazası
    ws = wb.create_sheet("Qida bazası")
    header(ws, 1, ["Məhsul (italyanca)", "İtalyanca", "Qrup", "kkal / 100 q", "Protein, q", "Karbohidrat, q", "Yağ, q", "Porsiya"],
           [48, 30, 16, 12, 11, 13, 9, 40])
    for i, f in enumerate(FOODS):
        r = 2 + i
        for j, v in enumerate([lab(f), f[2], f[3], f[4], f[5], f[6], f[7], f[8]], 1):
            ws.cell(row=r, column=j, value=v).font = N
    first_custom = 2 + len(FOODS)
    for r in range(first_custom, first_custom + 30):
        for j in range(1, 8): ws.cell(row=r, column=j).fill = YFILL
        ws.cell(row=r, column=3, value="Öz məhsulum" if r == first_custom else None)
    ws.cell(row=first_custom, column=1).comment = Comment("Öz məhsulunu bu sarı sətirlərə yaz: ad + per 100 g dəyərləri.", "Plan")
    FOOD_LAST = first_custom + 29
    ws.freeze_panes = "A2"

    # ---------------- Gündəlik
    ws = wb.create_sheet("Gündəlik")
    header(ws, 1, ["Tarix", "Yemək", "Məhsul", "Qram", "kkal", "Protein", "Karb.", "Yağ", "Qeyd"],
           [12, 16, 46, 9, 9, 9, 9, 9, 30])
    ws.freeze_panes = "A2"
    dv_meal = DataValidation(type="list", formula1='"' + ",".join(MEALS) + '"', allow_blank=True)
    dv_food = DataValidation(type="list", formula1=f"='Qida bazası'!$A$2:$A${FOOD_LAST}", allow_blank=True)
    ws.add_data_validation(dv_meal); ws.add_data_validation(dv_food)
    dv_meal.add(f"B2:B{LOGROWS+1}"); dv_food.add(f"C2:C{LOGROWS+1}")
    # example rows: tomorrow (Friday) menu
    fri = [d for d in MENU if d[0] == "Cümə"][0]
    ex = []
    for meal, title, items in fri[2]:
        for fid, g in items: ex.append((START, meal, lab(FOOD[fid]), g))
    rng = f"'Qida bazası'!$A$2:$A${FOOD_LAST}"
    for r in range(2, LOGROWS + 2):
        if r - 2 < len(ex):
            d, meal, name, g = ex[r - 2]
            for col, v in ((1, d), (2, meal), (3, name), (4, g)):
                c = ws.cell(row=r, column=col, value=v); c.font = INP
            ws.cell(row=r, column=9, value="nümunə — dəyiş və ya sil" if r == 2 else None).font = MUT
        ws.cell(row=r, column=1).number_format = "DD.MM.YYYY"
        for k, col in ((5, "D"), (6, "E"), (7, "F"), (8, "G")):
            f = (f'=IF(OR(C{r}="",D{r}=""),"",IFERROR(INDEX(\'Qida bazası\'!${col}$2:${col}${FOOD_LAST},'
                 f'MATCH(C{r},{rng},0))*D{r}/100,"?"))')
            c = ws.cell(row=r, column=k, value=f); c.font = N; c.number_format = "0" if k == 5 else "0.0"
        for k in range(1, 5): ws.cell(row=r, column=k).fill = YFILL if r - 2 >= len(ex) and r < 60 else PatternFill()

    # ---------------- Günlər (daily analysis)
    ws = wb.create_sheet("Günlər")
    header(ws, 1, ["Tarix", "Gün", "Tip", "Tennis, dəq", "Tennis növü", "Norma kkal", "Yeyilən kkal", "Fərq", "Status",
                   "Protein", "Protein min", "Karb.", "Yağ", "Çəki, kq", "7 gün orta çəki"],
           [12, 16, 11, 10, 11, 11, 11, 9, 11, 9, 11, 9, 9, 10, 13])
    ws.freeze_panes = "B2"
    dv_type = DataValidation(type="list", formula1='"Zal,İstirahət"', allow_blank=False)
    dv_tk = DataValidation(type="list", formula1='"Tək,Cüt"', allow_blank=True)
    ws.add_data_validation(dv_type); ws.add_data_validation(dv_tk)
    last = NDAYS + 1
    dv_type.add(f"C2:C{last}"); dv_tk.add(f"E2:E{last}")
    az_days = ["Bazar ertəsi", "Çərşənbə axşamı", "Çərşənbə", "Cümə axşamı", "Cümə", "Şənbə", "Bazar"]
    LOG = f"Gündəlik!$A$2:$A${LOGROWS+1}"
    for i in range(NDAYS):
        r = 2 + i; d = START + datetime.timedelta(days=i)
        ws.cell(row=r, column=1, value=d).number_format = "DD.MM.YYYY"
        ws.cell(row=r, column=2, value=az_days[d.weekday()])
        c = ws.cell(row=r, column=3, value="Zal" if d.weekday() in (0, 2, 4) else "İstirahət"); c.font = INP
        c = ws.cell(row=r, column=4); c.font = INP; c.fill = YFILL
        c = ws.cell(row=r, column=5, value="Tək"); c.font = INP
        ws.cell(row=r, column=6, value=(f'=ROUND(((10*Profil!$B$4+6.25*Profil!$B$5-5*Profil!$B$6+5)*Profil!$B$7+Profil!$B$8'
                                         f'+IF(C{r}="Zal",Profil!$B$9,0)+IF(N(D{r})>0,IF(E{r}="Cüt",5,7)*Profil!$B$4*D{r}/60,0))/10,0)*10')).number_format = "#,##0"
        ws.cell(row=r, column=7, value=f'=IF(COUNTIF({LOG},A{r})=0,"",SUMIFS(Gündəlik!$E$2:$E${LOGROWS+1},{LOG},A{r}))').number_format = "#,##0"
        ws.cell(row=r, column=8, value=f'=IF(G{r}="","",G{r}-F{r})').number_format = "+#,##0;-#,##0;0"
        ws.cell(row=r, column=9, value=f'=IF(G{r}="","",IF(G{r}<F{r}*0.9,"az",IF(G{r}>F{r}*1.15,"çox","normada")))')
        ws.cell(row=r, column=10, value=f'=IF(G{r}="","",SUMIFS(Gündəlik!$F$2:$F${LOGROWS+1},{LOG},A{r}))').number_format = "0"
        ws.cell(row=r, column=11, value="=ROUND(Profil!$B$10*Profil!$B$4,0)").number_format = "0"
        ws.cell(row=r, column=12, value=f'=IF(G{r}="","",SUMIFS(Gündəlik!$G$2:$G${LOGROWS+1},{LOG},A{r}))').number_format = "0"
        ws.cell(row=r, column=13, value=f'=IF(G{r}="","",SUMIFS(Gündəlik!$H$2:$H${LOGROWS+1},{LOG},A{r}))').number_format = "0"
        c = ws.cell(row=r, column=14); c.font = INP; c.fill = YFILL; c.number_format = "0.0"
        ws.cell(row=r, column=15, value=(f'=IF(COUNTIFS($A$2:$A${last},">="&(A{r}-6),$A$2:$A${last},"<="&A{r},$N$2:$N${last},">0")<3,"",'
                                          f'AVERAGEIFS($N$2:$N${last},$A$2:$A${last},">="&(A{r}-6),$A$2:$A${last},"<="&A{r},$N$2:$N${last},">0"))')).number_format = "0.00"
    green = PatternFill("solid", fgColor="D9EADF"); amber = PatternFill("solid", fgColor="FBE3B5"); red = PatternFill("solid", fgColor="F6CFCB")
    ws.conditional_formatting.add(f"I2:I{last}", CellIsRule(operator="equal", formula=['"normada"'], fill=green))
    ws.conditional_formatting.add(f"I2:I{last}", CellIsRule(operator="equal", formula=['"az"'], fill=amber))
    ws.conditional_formatting.add(f"I2:I{last}", CellIsRule(operator="equal", formula=['"çox"'], fill=red))
    ws.conditional_formatting.add(f"J2:J{last}", FormulaRule(formula=[f'AND(J2<>"",J2<K2*0.95)'], fill=amber))
    ch = BarChart(); ch.type = "col"; ch.title = "Yeyilən kkal və norma (ilk 28 gün)"; ch.height = 8; ch.width = 22
    ch.add_data(Reference(ws, min_col=7, min_row=1, max_row=29), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=29))
    ln = LineChart(); ln.add_data(Reference(ws, min_col=6, min_row=1, max_row=29), titles_from_data=True)
    ch += ln
    ws.add_chart(ch, "Q2")

    # ---------------- Həftəlik
    ws = wb.create_sheet("Həftəlik")
    header(ws, 1, ["Həftə (B.e.)", "Yazılmış gün", "Orta kkal", "Orta norma", "Orta protein", "Zal günləri", "Tennis, dəq",
                   "Orta çəki", "Dəyişmə, kq", "Tövsiyə"], [13, 12, 11, 11, 12, 11, 11, 10, 11, 46])
    mon0 = START - datetime.timedelta(days=START.weekday())
    weeks = NDAYS // 7 + 1
    A = f"Günlər!$A$2:$A${last}"
    for w in range(weeks):
        r = 2 + w
        ws.cell(row=r, column=1, value=mon0 + datetime.timedelta(days=7 * w)).number_format = "DD.MM.YYYY"
        cond = f'{A},">="&A{r},{A},"<"&(A{r}+7)'
        ws.cell(row=r, column=2, value=f'=COUNTIFS({cond},Günlər!$G$2:$G${last},">0")')
        ws.cell(row=r, column=3, value=f'=IF(B{r}=0,"",AVERAGEIFS(Günlər!$G$2:$G${last},{cond},Günlər!$G$2:$G${last},">0"))').number_format = "#,##0"
        ws.cell(row=r, column=4, value=f'=IF(B{r}=0,"",AVERAGEIFS(Günlər!$F$2:$F${last},{cond},Günlər!$G$2:$G${last},">0"))').number_format = "#,##0"
        ws.cell(row=r, column=5, value=f'=IF(B{r}=0,"",AVERAGEIFS(Günlər!$J$2:$J${last},{cond},Günlər!$G$2:$G${last},">0"))').number_format = "0"
        ws.cell(row=r, column=6, value=f'=COUNTIFS({cond},Günlər!$C$2:$C${last},"Zal")')
        ws.cell(row=r, column=7, value=f'=SUMIFS(Günlər!$D$2:$D${last},{cond})')
        ws.cell(row=r, column=8, value=f'=IF(COUNTIFS({cond},Günlər!$N$2:$N${last},">0")<3,"",AVERAGEIFS(Günlər!$N$2:$N${last},{cond},Günlər!$N$2:$N${last},">0"))').number_format = "0.00"
        if w == 0:
            ws.cell(row=r, column=9, value="")
        else:
            ws.cell(row=r, column=9, value=f'=IF(OR(H{r}="",H{r-1}=""),"",H{r}-H{r-1})').number_format = "+0.00;-0.00;0.00"
        ws.cell(row=r, column=10, value=(f'=IF(B{r}=0,"",IF(AND(I{r}<>"",I{r}<0.1),"Artım yavaşdır: Profil → artıq kaloriyə +150",'
                                          f'IF(AND(I{r}<>"",I{r}>0.4),"Çox sürətli: Profil → artıq kaloridən −150",'
                                          f'IF(C{r}<D{r}*0.9,"Az yeyirsən: hər gün +1 qəlyanaltı",IF(E{r}<Profil!$B$20*0.95,"Proteini artır","Əla, belə davam et")))))'))
    ws["J1"].comment = Comment("Çəki dəyişməsi üçün hər həftə ən azı 3 çəki lazımdır.", "Plan")

    # ---------------- Menyu
    ws = wb.create_sheet("Menyu")
    header(ws, 1, ["Gün", "Tip", "Yemək", "Saat", "Nə", "Məhsul", "Qram", "kkal", "Protein", "Karb.", "Yağ"],
           [16, 10, 15, 7, 38, 44, 7, 8, 8, 8, 8])
    r = 2
    times = [t for t, _, _ in TIMING["rest"][1]]
    for day, typ, meals in MENU:
        first = r
        for mi, (meal, title, items) in enumerate(meals):
            for ii, (fid, g) in enumerate(items):
                f = FOOD[fid]
                vals = [day if r == first else None, ("Zal" if typ == "gym" else "İstirahət") if r == first else None,
                        meal if ii == 0 else None, times[mi] if ii == 0 else None, title if ii == 0 else None, lab(f), g,
                        f"=ROUND(G{r}*INDEX('Qida bazası'!$D:$D,MATCH(F{r},'Qida bazası'!$A:$A,0))/100,0)",
                        f"=ROUND(G{r}*INDEX('Qida bazası'!$E:$E,MATCH(F{r},'Qida bazası'!$A:$A,0))/100,1)",
                        f"=ROUND(G{r}*INDEX('Qida bazası'!$F:$F,MATCH(F{r},'Qida bazası'!$A:$A,0))/100,1)",
                        f"=ROUND(G{r}*INDEX('Qida bazası'!$G:$G,MATCH(F{r},'Qida bazası'!$A:$A,0))/100,1)"]
                for j, v in enumerate(vals, 1):
                    c = ws.cell(row=r, column=j, value=v); c.font = B if j in (1, 3) else N
                r += 1
        ws.cell(row=r, column=6, value=f"Cəmi: {day}").font = B
        for j, col in ((8, "H"), (9, "I"), (10, "J"), (11, "K")):
            c = ws.cell(row=r, column=j, value=f"=SUM({col}{first}:{col}{r-1})"); c.font = B; c.fill = SFILL
        norm = "Profil!$B$17" if typ == "gym" else "Profil!$B$16"
        r += 1
        ws.cell(row=r, column=7, value="norma").font = MUT
        c = ws.cell(row=r, column=8, value=f"={norm}"); c.font = LINK; c.number_format = "#,##0"
        c = ws.cell(row=r, column=9, value="=Profil!$B$20"); c.font = LINK
        r += 2
    ws.freeze_panes = "A2"

    # ---------------- Yemək saatları
    ws = wb.create_sheet("Yemək saatları")
    ws.column_dimensions["A"].width = 9; ws.column_dimensions["B"].width = 34; ws.column_dimensions["C"].width = 60
    r = 1
    for k in ("rest", "morning", "midday", "evening"):
        title, rows = TIMING[k]
        ws.cell(row=r, column=1, value=title).font = H1; r += 1
        for t, m, n in rows:
            ws.cell(row=r, column=1, value=t).font = B if m == "MƏŞQ" else N
            ws.cell(row=r, column=2, value=m).font = B if m == "MƏŞQ" else N
            ws.cell(row=r, column=3, value=n).font = MUT
            r += 1
        r += 1
    ws.cell(row=r, column=1, value="Tennis günü").font = H1; r += 1
    for t in TENNIS_TIPS:
        c = ws.cell(row=r, column=1, value="• " + t); c.font = N; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c.alignment = Alignment(wrap_text=True); ws.row_dimensions[r].height = 30; r += 1
    r += 1
    ws.cell(row=r, column=1, value="Əsas qaydalar").font = H1; r += 1
    for t in GENERAL_RULES:
        c = ws.cell(row=r, column=1, value="• " + t); c.font = N; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c.alignment = Alignment(wrap_text=True); ws.row_dimensions[r].height = 30; r += 1

    # ---------------- Məşq
    ws = wb.create_sheet("Məşq")
    header(ws, 1, ["Məşq", "Hərəkət", "Set × təkrar", "Fasilə", "Texnika", "Animasiya (repo)"], [7, 36, 16, 10, 70, 30])
    r = 2
    for k in ("A", "B"):
        for i, n, s, rest, cue in WORKOUT[k]:
            vals = [k, n, s, rest, cue, ex_db[i]["gif_url"]]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v); c.font = N; c.alignment = WRAP
            r += 1
        r += 1
    ws.cell(row=r, column=1, value="Həftəlik cədvəl").font = H1; r += 1
    for d, t in WEEK_SCHEDULE:
        ws.cell(row=r, column=2, value=d).font = B; ws.cell(row=r, column=3, value=t).font = N; r += 1
    r += 1
    ws.cell(row=r, column=1, value="Qaydalar").font = H1; r += 1
    for t in TRAINING_RULES:
        c = ws.cell(row=r, column=2, value="• " + t); c.font = N; ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c.alignment = Alignment(wrap_text=True); ws.row_dimensions[r].height = 28; r += 1
    r += 1
    ws.cell(row=r, column=1, value="Məşq jurnalı").font = H1; r += 1
    jh = ["Tarix", "Hərəkət", "Set 1 (kq × təkrar)", "Set 2", "Set 3", "Qeyd"]
    for j, l in enumerate(jh, 1):
        c = ws.cell(row=r, column=j, value=l); c.font = HB; c.fill = HFILL
    r += 1
    for _ in range(40):
        for j in range(1, 7): ws.cell(row=r, column=j).fill = YFILL
        r += 1

    # ---------------- Alış-veriş
    ws = wb.create_sheet("Alış-veriş")
    header(ws, 1, ["Məhsul", "İtalyanca", "Həftəlik, q", "Miqdar", "≈ €/kq", "≈ €"], [30, 30, 12, 14, 9, 9])
    for i, s in enumerate(SHOP):
        r = 2 + i
        for j, v in enumerate([s["name"], s["it"], s["grams"], s["qty"]], 1): ws.cell(row=r, column=j, value=v).font = N
        c = ws.cell(row=r, column=5, value=s["eur_kg"]); c.font = INP; c.number_format = "0.00"
        ws.cell(row=r, column=6, value=f"=C{r}/1000*E{r}").number_format = "0.00"
    r = 2 + len(SHOP)
    ws.cell(row=r, column=1, value="Cəmi (həftə)").font = B
    c = ws.cell(row=r, column=6, value=f"=SUM(F2:F{r-1})"); c.font = B; c.number_format = "0.00"
    ws.cell(row=r + 1, column=1, value="Qiymətlər Bolonya supermarketləri üçün təxminidir (2026). Bar/restoran yeməkləri daxil deyil.").font = MUT
    ws["E1"].comment = Comment("Təxmini supermarket qiyməti (Lidl/Conad/Esselunga). Öz çekinə görə dəyiş.", "Plan")

    # ---------------- Bolonya
    ws = wb.create_sheet("Bolonya")
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 100
    ws["A1"] = "Bolonyada necə yemək"; ws["A1"].font = H1
    for i, (h, t) in enumerate(BOLOGNA_TIPS):
        r = 3 + i
        ws.cell(row=r, column=1, value=h).font = B
        c = ws.cell(row=r, column=2, value=t); c.font = N; c.alignment = WRAP
        ws.row_dimensions[r].height = 45
        ws.cell(row=r, column=1).alignment = WRAP

    for s in wb.worksheets:
        for row in s.iter_rows():
            for c in row:
                if c.value is not None and (c.font is None or c.font.name != F):
                    c.font = Font(name=F, size=10)
    wb.active = 0
    out = os.path.join(OUT_PLAN, "qidalanma-plani.xlsx")
    wb.save(out)
    return out

if __name__ == "__main__":
    build_html()
    print(build_xlsx())
    t = targets(); print(t)
    print("shopping €", round(sum(s["eur"] for s in SHOP), 2))
